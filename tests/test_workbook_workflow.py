from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook

from src.planning import create_recommendations
from src.workbook import evaluate_cell, inspect_stock_workbook, write_stock_output


def make_test_workbook(path: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "CAT Plan Jan-26"
    ws["A5"] = "D123 PIPELINE PLAN -261-1544"
    ws["A6"] = "C9 CRANKSHAFT"
    ws["A7"] = "Activity"
    for col, month in zip((2, 3, 4), (datetime(2026, 1, 1), datetime(2026, 2, 1), datetime(2026, 3, 1))):
        ws.cell(7, col).value = month
    ws["A8"] = "Shipment ex-BFL (Nos)"
    ws["B8"], ws["C8"], ws["D8"] = 100, 120, 130
    ws["A9"] = "Arrivals at Warehouse"
    ws["B9"], ws["C9"], ws["D9"] = 40, 50, 60
    ws["A10"] = "Forecasted consumption at CAT"
    ws["B10"], ws["C10"], ws["D10"] = 80, 90, 100
    ws["A11"] = "Effective Stock in Warehouse"
    ws["B11"] = 500
    ws["C11"] = "=B11+C9-C10"
    ws["D11"] = "=C11+D9-D10"
    ws["A12"] = "Stock in Transit"
    ws["B12"], ws["C12"], ws["D12"] = 200, 210, 220
    ws["A13"] = "Stock in Pipeline"
    ws["B13"], ws["C13"], ws["D13"] = 700, 800, 900
    ws["A14"] = "Inventory level in months"
    ws["B14"], ws["C14"], ws["D14"] = 5.5, 4.4, 3.3
    wb.save(path)
    return path


def test_detects_stock_section_and_month_headers(tmp_path):
    path = make_test_workbook(tmp_path / "stock.xlsx")
    info = inspect_stock_workbook(str(path), "CAT Plan Jan-26")
    assert len(info["sections"]) == 1
    section = info["sections"][0]
    assert section.part_number == "261-1544"
    assert section.shipment_row == 8
    assert [m.label for m in section.months] == ["Jan-26", "Feb-26", "Mar-26"]
    assert section.rows["pipeline"] == 13


def test_safe_formula_evaluator_supports_overrides(tmp_path):
    path = make_test_workbook(tmp_path / "stock.xlsx")
    wb = load_workbook(path, data_only=False)
    ws = wb["CAT Plan Jan-26"]
    actual = evaluate_cell(ws, ws, "C11", overrides={"B11": 600})
    assert actual == 560  # 600 + arrivals(50) - consumption(90)
    wb.close()


def test_export_writes_shipment_without_replacing_formula(tmp_path):
    path = make_test_workbook(tmp_path / "stock.xlsx")
    info = inspect_stock_workbook(str(path), "CAT Plan Jan-26")
    section = info["sections"][0]
    out = tmp_path / "final.xlsx"
    write_stock_output(path, "CAT Plan Jan-26", section, {"Feb-26": 777}, out, demand_quantities={"Feb-26": 222})
    wb = load_workbook(out, data_only=False)
    ws = wb["CAT Plan Jan-26"]
    assert ws["C8"].value == 777
    assert ws["C10"].value == 222
    assert ws["C11"].value == "=B11+C9-C10"
    assert ws["D8"].value == 130
    wb.close()


def test_recommendation_plan_has_editable_approved_quantity(tmp_path):
    path = make_test_workbook(tmp_path / "stock.xlsx")
    section = inspect_stock_workbook(str(path), "CAT Plan Jan-26")["sections"][0]
    df = create_recommendations(
        str(path), "CAT Plan Jan-26", section, "Feb-26",
        {"transit_months": 1, "target_coverage_months": 1, "rounding_multiple": 10},
        horizon_months=2,
    )
    assert list(df["Month"]) == ["Feb-26", "Mar-26"]
    assert "Suggested shipment" in df.columns
    assert "Approved shipment" in df.columns
    assert (df["Approved shipment"] >= 0).all()


def test_wide_maniana_demand_mapping(tmp_path):
    from src.demand import extract_demand_by_part

    path = tmp_path / "maniana.csv"
    path.write_text("Part Number,Jan-26,Feb-26,Mar-26\n261-1544,100,120,140\n282-7958,10,20,30\n261-1544,5,5,5\n", encoding="utf-8")
    mapping = {
        "sheet_name": "CSV", "header_row": 1, "layout": "wide", "part_column_index": 0,
        "month_columns": [
            {"label": "Jan-26", "index": 1},
            {"label": "Feb-26", "index": 2},
            {"label": "Mar-26", "index": 3},
        ],
    }
    demand, matches = extract_demand_by_part(path, mapping, "261-1544")
    assert matches == 2
    assert demand == {"Jan-26": 105.0, "Feb-26": 125.0, "Mar-26": 145.0}


def test_long_maniana_demand_mapping(tmp_path):
    from src.demand import extract_demand_by_part

    path = tmp_path / "maniana.csv"
    path.write_text("Part,Date,Demand\n261-1544,20260101,100\n261-1544,20260201,125\n282-7958,20260201,80\n", encoding="utf-8")
    mapping = {
        "sheet_name": "CSV", "header_row": 1, "layout": "long", "part_column_index": 0,
        "date_column_index": 1, "quantity_column_index": 2,
    }
    demand, matches = extract_demand_by_part(path, mapping, "261-1544")
    assert matches == 2
    assert demand == {"Jan-26": 100.0, "Feb-26": 125.0}


def test_revision_history_preserves_demand_snapshot(tmp_path, monkeypatch):
    import src.history as history

    monkeypatch.setattr(history, "DB_PATH", tmp_path / "history.sqlite")
    revision_id = history.save_revision(
        workbook_name="stock.xlsx", sheet_name="CAT Plan Jan-26", section_title="261-1544",
        part_number="261-1544", planning_month="Feb-26", revision_type="Manufacturing Update — 7th",
        status="Draft", notes="manufacturing update", quantities={"Feb-26": 120},
        demand_overrides={"Feb-26": 123.0},
    )
    revision = history.get_revision(revision_id)
    assert history.quantities_from_revision(revision) == {"Feb-26": 120}
    assert revision["demand_overrides"] == {"Feb-26": 123.0}
