"""Workbook inspection, safe formula evaluation, and template-preserving export."""
from __future__ import annotations

import ast
import re
from functools import lru_cache
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

LABEL_ALIASES = {
    "shipment": ("shipment ex-bfl", "shipment ex bfl"),
    "arrivals": ("arrivals at warehouse",),
    "consumption": ("forecasted consumption at cat", "forecasted consumption"),
    "effective_stock": ("effective stock in warehouse",),
    "transit": ("stock in transit",),
    "pipeline": ("stock in pipeline",),
    "coverage": ("inventory level in months", "inventory level"),
}


@dataclass(frozen=True)
class MonthColumn:
    column: int
    label: str
    period: date


@dataclass
class StockSection:
    section_id: str
    title: str
    part_number: str
    header_row: int
    shipment_row: int
    rows: dict[str, int]
    months: list[MonthColumn]


def _normalise(value: Any) -> str:
    return re.sub(r"\s+", " ", str(value or "")).strip().lower()


def _parse_date(value: Any) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str):
        text = value.strip()
        for fmt in ("%b-%y", "%B-%y", "%Y-%m", "%Y-%m-%d", "%d-%b-%y", "%b %Y"):
            try:
                return datetime.strptime(text, fmt).date()
            except ValueError:
                pass
    return None


def month_label(period: date) -> str:
    return period.strftime("%b-%y")


def section_part_number(title: str) -> str:
    matches = re.findall(r"\b\d{3,4}-\d{4}\b", title or "")
    return matches[-1] if matches else "Unmapped part"


def _label_kind(value: Any) -> str | None:
    text = _normalise(value)
    for key, aliases in LABEL_ALIASES.items():
        if any(text.startswith(alias) for alias in aliases):
            return key
    return None


def _find_dates(ws, row_number: int) -> list[MonthColumn]:
    months: list[MonthColumn] = []
    for col in range(2, ws.max_column + 1):
        value = ws.cell(row_number, col).value
        parsed = _parse_date(value)
        if parsed:
            months.append(MonthColumn(col, month_label(parsed), parsed))
    return months


@lru_cache(maxsize=48)
def inspect_stock_workbook(path: str | Path, sheet_name: str | None = None) -> dict[str, Any]:
    """Return basic metadata and discover stock-plan sections in a workbook."""
    path = Path(path)
    if not path.exists():
        raise FileNotFoundError(f"Workbook not found: {path}")
    wb = load_workbook(path, data_only=False, read_only=False)
    sheet_names = wb.sheetnames
    selected = sheet_name if sheet_name in sheet_names else (sheet_names[0] if sheet_names else None)
    sections: list[StockSection] = []
    if selected:
        ws = wb[selected]
        for row in range(1, ws.max_row + 1):
            if _label_kind(ws.cell(row, 1).value) != "shipment":
                continue
            header_row = None
            months: list[MonthColumn] = []
            for candidate in range(max(1, row - 1), max(0, row - 9), -1):
                possible = _find_dates(ws, candidate)
                if possible and (_normalise(ws.cell(candidate, 1).value) == "activity" or len(possible) >= 3):
                    header_row, months = candidate, possible
                    break
            if not header_row or len(months) < 3:
                continue
            title = ""
            for candidate in range(header_row - 1, max(0, header_row - 7), -1):
                value = ws.cell(candidate, 1).value
                if value and "pipeline plan" in _normalise(value):
                    title = str(value).strip()
                    break
            if not title:
                title = f"Stock plan section at row {row}"
            rows: dict[str, int] = {"shipment": row}
            for next_row in range(row + 1, min(ws.max_row, row + 10) + 1):
                val = ws.cell(next_row, 1).value
                if _normalise(val) == "activity" or (val and "pipeline plan" in _normalise(val)):
                    break
                kind = _label_kind(val)
                if kind and kind not in rows:
                    rows[kind] = next_row
            part = section_part_number(title)
            section_id = f"{selected}::{row}::{part}"
            sections.append(StockSection(section_id, title, part, header_row, row, rows, months))
    wb.close()
    return {
        "path": str(path),
        "filename": path.name,
        "sheets": sheet_names,
        "selected_sheet": selected,
        "sections": sections,
        "size_bytes": path.stat().st_size,
    }


def get_section(path: str | Path, sheet_name: str, section_id: str) -> StockSection:
    result = inspect_stock_workbook(path, sheet_name)
    for section in result["sections"]:
        if section.section_id == section_id:
            return section
    raise ValueError("The selected stock-plan section could not be found. Please reselect the worksheet and section.")


def _safe_eval(node: ast.AST) -> float:
    if isinstance(node, ast.Expression):
        return _safe_eval(node.body)
    if isinstance(node, ast.Constant) and isinstance(node.value, (int, float)):
        return float(node.value)
    if isinstance(node, ast.UnaryOp) and isinstance(node.op, (ast.UAdd, ast.USub)):
        val = _safe_eval(node.operand)
        return val if isinstance(node.op, ast.UAdd) else -val
    if isinstance(node, ast.BinOp) and isinstance(node.op, (ast.Add, ast.Sub, ast.Mult, ast.Div, ast.Pow, ast.Mod)):
        left, right = _safe_eval(node.left), _safe_eval(node.right)
        if isinstance(node.op, ast.Add): return left + right
        if isinstance(node.op, ast.Sub): return left - right
        if isinstance(node.op, ast.Mult): return left * right
        if isinstance(node.op, ast.Div): return left / right if right else 0.0
        if isinstance(node.op, ast.Pow): return left ** right
        if isinstance(node.op, ast.Mod): return left % right if right else 0.0
    raise ValueError("Unsupported formula expression")


_CELL_REF = re.compile(r"(?<![A-Z0-9_])\$?([A-Z]{1,3})\$?(\d+)(?![A-Z0-9_])", re.I)


def _numeric(value: Any) -> float:
    if value is None or value == "":
        return 0.0
    if isinstance(value, bool):
        return float(value)
    if isinstance(value, (int, float)):
        return float(value)
    if isinstance(value, str):
        try:
            return float(value.replace(",", ""))
        except ValueError:
            return 0.0
    return 0.0


def evaluate_cell(ws, cached_ws, coordinate: str, overrides: dict[str, float] | None = None,
                  memo: dict[str, float] | None = None, stack: set[str] | None = None,
                  boundary_col: int | None = None) -> float:
    """Evaluate simple Excel arithmetic formulas with override support.

    This is intentionally a small evaluator for stock-plan formulas (cell references and
    arithmetic). Unsupported formulas fall back to Excel's cached result where available.
    It never executes arbitrary Python code.
    """
    overrides = overrides or {}
    memo = memo if memo is not None else {}
    stack = stack if stack is not None else set()
    coordinate = coordinate.upper().replace("$", "")
    if coordinate in overrides:
        return _numeric(overrides[coordinate])
    # Before the earliest overridden shipment month, cached Excel results are the stable boundary.
    # This prevents recalculating a decade of historical formulas for a short planning scenario.
    ref_col_letters = re.match(r"[A-Z]+", coordinate)
    if boundary_col is not None and ref_col_letters:
        from openpyxl.utils.cell import column_index_from_string
        if column_index_from_string(ref_col_letters.group(0)) < boundary_col:
            return _numeric(cached_ws[coordinate].value)
    if coordinate in memo:
        return memo[coordinate]
    if coordinate in stack:
        return _numeric(cached_ws[coordinate].value)
    stack.add(coordinate)
    value = ws[coordinate].value
    if isinstance(value, str) and value.startswith("="):
        formula = value[1:]
        variables: dict[str, float] = {}
        def replace_ref(match: re.Match[str]) -> str:
            ref = f"{match.group(1).upper()}{match.group(2)}"
            variable = f"REF{len(variables)}"
            variables[variable] = evaluate_cell(ws, cached_ws, ref, overrides, memo, stack, boundary_col)
            return variable
        candidate = _CELL_REF.sub(replace_ref, formula)
        try:
            tree = ast.parse(candidate, mode="eval")
            # Replace the safe reference variables with numeric constants.
            class RefsToConstants(ast.NodeTransformer):
                def visit_Name(self, node):
                    if node.id in variables:
                        return ast.copy_location(ast.Constant(value=variables[node.id]), node)
                    raise ValueError("Unsupported name in formula")
            tree = RefsToConstants().visit(tree)
            ast.fix_missing_locations(tree)
            result = _safe_eval(tree)
        except Exception:
            result = _numeric(cached_ws[coordinate].value)
    else:
        result = _numeric(value)
    stack.discard(coordinate)
    memo[coordinate] = result
    return result


@lru_cache(maxsize=8)
def _load_sheet_snapshot(path_text: str, file_mtime_ns: int, sheet_name: str) -> tuple[dict[str, Any], dict[str, Any]]:
    """Cache compact formula/cached-value maps so page reruns do not repeatedly open large workbooks."""
    formula_wb = load_workbook(path_text, data_only=False, read_only=False)
    cached_wb = load_workbook(path_text, data_only=True, read_only=False)
    if sheet_name not in formula_wb.sheetnames:
        formula_wb.close()
        cached_wb.close()
        raise ValueError(f"Worksheet '{sheet_name}' was not found in the workbook.")
    formula_ws, cached_ws = formula_wb[sheet_name], cached_wb[sheet_name]
    formulas: dict[str, Any] = {}
    cached: dict[str, Any] = {}
    for row in formula_ws.iter_rows():
        for cell in row:
            if cell.value is not None:
                formulas[cell.coordinate] = cell.value
    for row in cached_ws.iter_rows():
        for cell in row:
            if cell.value is not None:
                cached[cell.coordinate] = cell.value
    formula_wb.close()
    cached_wb.close()
    return formulas, cached


def _evaluate_snapshot(coordinate: str, formulas: dict[str, Any], cached: dict[str, Any],
                       overrides: dict[str, float], memo: dict[str, float], stack: set[str],
                       boundary_col: int | None = None) -> float:
    coordinate = coordinate.upper().replace("$", "")
    if coordinate in overrides:
        return _numeric(overrides[coordinate])
    match_col = re.match(r"[A-Z]+", coordinate)
    if boundary_col is not None and match_col:
        from openpyxl.utils.cell import column_index_from_string
        if column_index_from_string(match_col.group(0)) < boundary_col:
            return _numeric(cached.get(coordinate))
    if coordinate in memo:
        return memo[coordinate]
    if coordinate in stack:
        return _numeric(cached.get(coordinate))
    stack.add(coordinate)
    raw = formulas.get(coordinate)
    if isinstance(raw, str) and raw.startswith("="):
        variables: dict[str, float] = {}
        def replace_ref(match: re.Match[str]) -> str:
            ref = f"{match.group(1).upper()}{match.group(2)}"
            variable = f"REF{len(variables)}"
            variables[variable] = _evaluate_snapshot(ref, formulas, cached, overrides, memo, stack, boundary_col)
            return variable
        candidate = _CELL_REF.sub(replace_ref, raw[1:])
        try:
            tree = ast.parse(candidate, mode="eval")
            class RefsToConstants(ast.NodeTransformer):
                def visit_Name(self, node):
                    if node.id in variables:
                        return ast.copy_location(ast.Constant(value=variables[node.id]), node)
                    raise ValueError("Unsupported name in formula")
            tree = RefsToConstants().visit(tree)
            ast.fix_missing_locations(tree)
            result = _safe_eval(tree)
        except Exception:
            result = _numeric(cached.get(coordinate))
    else:
        result = _numeric(raw)
    stack.discard(coordinate)
    memo[coordinate] = result
    return result


def workbook_values(path: str | Path, sheet_name: str, section: StockSection,
                    shipment_overrides: dict[str, float] | None = None) -> dict[str, list[float]]:
    """Read and, when possible, recalculate the section rows for a preview.

    The compact sheet snapshot is cached by path, modification time and worksheet, so editing
    a shipment quantity does not reopen a large historical workbook on every Streamlit rerun.
    """
    path_obj = Path(path).resolve()
    shipment_overrides = {k.replace("$", "").upper(): v for k, v in (shipment_overrides or {}).items()}
    formulas, cached = _load_sheet_snapshot(str(path_obj), path_obj.stat().st_mtime_ns, sheet_name)
    row_values: dict[str, list[float]] = {}
    memo: dict[str, float] = {}
    from openpyxl.utils.cell import column_index_from_string
    override_cols = [column_index_from_string(re.match(r"[A-Z]+", c).group(0)) for c in shipment_overrides]
    boundary_col = min(override_cols) if override_cols else None
    for kind, row in section.rows.items():
        cells = [get_column_letter(m.column) + str(row) for m in section.months]
        if not shipment_overrides:
            row_values[kind] = [_numeric(cached.get(cell, formulas.get(cell))) for cell in cells]
        else:
            row_values[kind] = [
                _evaluate_snapshot(cell, formulas, cached, shipment_overrides, memo, set(), boundary_col)
                for cell in cells
            ]
    return row_values


def make_shipment_overrides(section: StockSection, quantities: dict[str, float]) -> dict[str, float]:
    """Convert month labels to actual worksheet cell references for active shipment quantities."""
    return {f"{get_column_letter(m.column)}{section.shipment_row}": float(quantities[m.label])
            for m in section.months if m.label in quantities and quantities[m.label] is not None}


def write_stock_output(source_path: str | Path, sheet_name: str, section: StockSection,
                       quantities: dict[str, float], output_path: str | Path,
                       demand_quantities: dict[str, float] | None = None) -> Path:
    """Write active shipment values into the existing template and force Excel recalculation.

    Formula cells, formatting, merged cells, and other worksheets are retained by openpyxl.
    Excel is asked to recalculate formulas when the exported workbook is opened.
    """
    source_path, output_path = Path(source_path), Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb = load_workbook(source_path, data_only=False, keep_links=True)
    if sheet_name not in wb.sheetnames:
        wb.close()
        raise ValueError(f"Worksheet '{sheet_name}' was not found in the source workbook.")
    ws = wb[sheet_name]
    for month in section.months:
        if month.label in quantities:
            raw = quantities[month.label]
            ws.cell(section.shipment_row, month.column).value = max(0, int(round(float(raw))))
    demand_quantities = demand_quantities or {}
    if "consumption" in section.rows:
        for month in section.months:
            if month.label in demand_quantities:
                ws.cell(section.rows["consumption"], month.column).value = max(0, int(round(float(demand_quantities[month.label]))))
    # Request recalculation when opened in desktop Excel or another compatible spreadsheet app.
    try:
        wb.calculation.fullCalcOnLoad = True
        wb.calculation.forceFullCalc = True
        wb.calculation.calcMode = "auto"
    except Exception:
        pass
    wb.save(output_path)
    wb.close()
    return output_path


def workbook_summary(path: str | Path) -> dict[str, Any]:
    meta = inspect_stock_workbook(path)
    return {"filename": meta["filename"], "sheets": len(meta["sheets"]),
            "size_mb": round(meta["size_bytes"] / (1024 * 1024), 2), "sheet_names": meta["sheets"]}
