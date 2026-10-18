"""Draft shipment recommendation logic. Company-specific parameters remain configurable."""
from __future__ import annotations

import math
from datetime import date
from typing import Any

import pandas as pd

from .workbook import StockSection, workbook_values


def _round_up(value: float, multiple: int) -> int:
    value = max(0, int(math.ceil(value)))
    multiple = max(1, int(multiple))
    return int(math.ceil(value / multiple) * multiple) if value else 0


def create_recommendations(path: str, sheet_name: str, section: StockSection,
                           planning_month: str, settings: dict[str, Any],
                           horizon_months: int = 18,
                           demand_overrides: dict[str, float] | None = None) -> pd.DataFrame:
    """Create a transparent provisional coverage-based shipment plan.

    This is intentionally labelled as a draft heuristic. Until Sales confirms its
    container/multiplier logic, users must review and approve every recommendation.
    """
    values = workbook_values(path, sheet_name, section)
    months = section.months
    start = next((i for i, m in enumerate(months) if m.label == planning_month), None)
    if start is None:
        raise ValueError(f"Planning month {planning_month} was not found in the selected worksheet.")
    end = min(len(months), start + max(1, int(horizon_months)))

    shipments = values.get("shipment", [0.0] * len(months))
    demand = values.get("consumption", [0.0] * len(months))
    demand_overrides = demand_overrides or {}
    for index, month in enumerate(months):
        if month.label in demand_overrides:
            demand[index] = max(0.0, float(demand_overrides[month.label]))
    pipeline = values.get("pipeline", [0.0] * len(months))
    transit_months = max(0, int(settings.get("transit_months", 2)))
    coverage_months = max(1, int(settings.get("target_coverage_months", 1)))
    rounding_multiple = max(1, int(settings.get("rounding_multiple", 1)))

    # Use the prior month's stock-in-pipeline as starting state, then project forward.
    pipeline_before = pipeline[start - 1] if start > 0 else 0.0
    rows: list[dict[str, Any]] = []
    for i in range(start, end):
        # A coverage target aims to cover demand through the transit horizon and retain
        # the configured number of demand months as a buffer.
        target_end = min(len(months), i + transit_months + coverage_months)
        period_demand = sum(max(0.0, float(demand[j] or 0)) for j in range(i, target_end))
        gap = max(0.0, period_demand - float(pipeline_before or 0))
        suggested = _round_up(gap, rounding_multiple)
        # Store a provisional scenario. Sales can replace this with its approved quantity.
        rows.append({
            "Month": months[i].label,
            "Customer demand": round(float(demand[i] or 0)),
            "Existing shipment": round(float(shipments[i] or 0)),
            "Pipeline before plan": round(float(pipeline_before or 0)),
            "Suggested shipment": suggested,
            "Approved shipment": suggested,
            "Worksheet column": months[i].column,
            "Month date": months[i].period,
        })
        pipeline_before = max(0.0, float(pipeline_before or 0) + suggested - float(demand[i] or 0))
    return pd.DataFrame(rows)


def active_quantities_from_df(df: pd.DataFrame) -> dict[str, int]:
    """Build the active month-to-quantity map from an editable planning table."""
    result: dict[str, int] = {}
    if df is None or df.empty:
        return result
    for _, row in df.iterrows():
        month = str(row["Month"])
        try:
            qty = max(0, int(round(float(row["Approved shipment"]))))
        except (TypeError, ValueError):
            qty = 0
        result[month] = qty
    return result


def create_scenario_preview(path: str, sheet_name: str, section: StockSection,
                            active_quantities: dict[str, int], starting_month: str,
                            demand_overrides: dict[str, float] | None = None) -> pd.DataFrame:
    """Project the familiar stock statement rows using the active shipment scenario."""
    import openpyxl
    from openpyxl.utils import get_column_letter
    from .workbook import evaluate_cell

    overrides = {f"{get_column_letter(m.column)}{section.shipment_row}": float(active_quantities[m.label])
                 for m in section.months if m.label in active_quantities}
    start = next((i for i, m in enumerate(section.months) if m.label == starting_month), 0)
    demand_overrides = demand_overrides or {}
    if "consumption" in section.rows:
        demand_row = section.rows["consumption"]
        for i, month in enumerate(section.months):
            if i >= start and month.label in demand_overrides:
                overrides[f"{get_column_letter(month.column)}{demand_row}"] = float(demand_overrides[month.label])
    values = workbook_values(path, sheet_name, section, overrides)
    end = min(len(section.months), start + 18)
    selected_months = section.months[start:end]
    rows: list[dict[str, Any]] = []

    # workbook_values performs simple formula-aware evaluation; formulas that are not
    # supported use the cached result from the source workbook until Excel recalculates.
    labels = [
        ("Shipment ex-BFL (Nos)", "shipment"),
        ("Arrivals at Warehouse", "arrivals"),
        ("Forecasted consumption at CAT", "consumption"),
        ("Effective Stock in Warehouse", "effective_stock"),
        ("Stock in Transit", "transit"),
        ("Stock in Pipeline", "pipeline"),
        ("Inventory level in months", "coverage"),
    ]
    for display, key in labels:
        if key not in values:
            continue
        row: dict[str, Any] = {"Activity": display}
        for i in range(start, end):
            qty = values[key][i]
            row[section.months[i].label] = round(qty, 2) if key == "coverage" else round(qty)
        rows.append(row)
    return pd.DataFrame(rows)
